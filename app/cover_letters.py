"""Extract plain-text cover letters from DOCX or PDF uploads."""

from __future__ import annotations

import asyncio
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from pypdf import PdfReader
from starlette.concurrency import run_in_threadpool


router = APIRouter(tags=["Cover letters"])

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TEXT_LENGTH = 60000
MAX_PDF_PAGES = 20
MAX_DOCX_EXPANDED_BYTES = 20 * 1024 * 1024

EXTRACT_LIMIT = asyncio.Semaphore(2)


def require_extract_access(request: Request) -> None:
    # Read the shared flag at request time to avoid a router import cycle.
    from app import jobs
    from app.auth_routes import _require_csrf, current_user

    if jobs.MULTI_USER_ENABLED:
        current_user(request)
        _require_csrf(request)


def clean_letter_text(value: str) -> str:
    value = value.replace("\x00", "")
    value = value.replace("\r\n", "\n").replace("\r", "\n")

    lines = [
        line.rstrip()
        for line in value.splitlines()
    ]

    return "\n".join(lines).strip()


def read_docx(data: bytes) -> str:
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()

            if len(entries) > 1000:
                raise ValueError("The Word document contains too many parts.")

            expanded_size = sum(
                item.file_size for item in entries
            )

            if expanded_size > MAX_DOCX_EXPANDED_BYTES:
                raise ValueError(
                    "The expanded Word document is too large."
                )

            if any(item.flag_bits & 1 for item in entries):
                raise ValueError(
                    "Encrypted Word documents are not supported."
                )

            if "word/document.xml" not in archive.namelist():
                raise ValueError("This is not a valid DOCX document.")

    except BadZipFile as exc:
        raise ValueError("This is not a valid DOCX document.") from exc

    document = Document(BytesIO(data))

    parts = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                parts.append(" | ".join(cells))

    return clean_letter_text("\n\n".join(parts))


def read_pdf(data: bytes) -> str:
    if b"%PDF-" not in data[:1024]:
        raise ValueError("This is not a valid PDF document.")

    reader = PdfReader(BytesIO(data), strict=False)

    if reader.is_encrypted:
        raise ValueError(
            "Encrypted PDFs are not supported. "
            "Please upload an unencrypted copy or paste the text."
        )

    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError(
            f"Please upload a PDF with no more than {MAX_PDF_PAGES} pages."
        )

    parts = []
    length = 0

    for page in reader.pages:
        text = page.extract_text() or ""
        length += len(text)

        if length > MAX_TEXT_LENGTH:
            raise ValueError(
                "The extracted letter exceeds 60,000 characters."
            )

        parts.append(text)

    return clean_letter_text("\n\n".join(parts))


def extract_letter(data: bytes, extension: str) -> str:
    if extension == ".docx":
        text = read_docx(data)
    elif extension == ".pdf":
        text = read_pdf(data)
    else:
        raise ValueError("Only DOCX and PDF files are supported.")

    if not text:
        raise ValueError(
            "No readable text was found. "
            "A scanned PDF may require OCR. "
            "Please paste the cover letter text instead."
        )

    if len(text) > MAX_TEXT_LENGTH:
        raise ValueError(
            "The extracted letter exceeds 60,000 characters."
        )

    return text


@router.post("/cover-letter/extract", dependencies=[Depends(require_extract_access)])
async def extract_cover_letter(
    file: UploadFile = File(...),
):
    filename = file.filename or ""
    extension = Path(filename).suffix.lower()

    try:
        if extension not in {".docx", ".pdf"}:
            raise HTTPException(
                status_code=422,
                detail="Please upload a DOCX or PDF file.",
            )

        data = await file.read(MAX_FILE_BYTES + 1)

        if not data:
            raise HTTPException(
                status_code=422,
                detail="The uploaded file is empty.",
            )

        if len(data) > MAX_FILE_BYTES:
            raise HTTPException(
                status_code=413,
                detail="The maximum upload size is 5 MB.",
            )

        async with EXTRACT_LIMIT:
            try:
                text = await run_in_threadpool(
                    extract_letter,
                    data,
                    extension,
                )
            except ValueError as exc:
                raise HTTPException(
                    status_code=422,
                    detail=str(exc),
                ) from exc
            except Exception as exc:
                raise HTTPException(
                    status_code=422,
                    detail=(
                        "The document could not be read. "
                        "Please check the file or paste the text."
                    ),
                ) from exc

        return {
            "text": text,
            "characters": len(text),
            "filename": filename,
            "saved": False,
        }

    finally:
        await file.close()
